from pathlib import Path
from typing import Optional
import click
from rich.console import Console
from rich.table import Table
from rich.progress import Progress, SpinnerColumn, TextColumn

from ..config import AppConfig
from ..state import AppState, run_batch
from ..models import Material, ProcessStatus, MaterialType
from ..utils import human_readable_size, human_readable_duration

console = Console()


def inspect_image(filepath: str) -> dict:
    try:
        from PIL import Image, ImageStat
        with Image.open(filepath) as img:
            width, height = img.size
            stat = ImageStat.Stat(img)
            if img.mode in ("L", "RGB", "RGBA"):
                brightness = sum(stat.mean[:3]) / len(stat.mean[:3]) / 255 * 100
            else:
                brightness = 50.0
            sharpness = 0.0
            try:
                gray = img.convert("L")
                edges = gray.filter(Image.Filter.Kernel((3, 3), (-1, -1, -1, -1, 8, -1, -1, -1, -1), 1, 0))
                edge_stat = ImageStat.Stat(edges)
                sharpness = edge_stat.mean[0] / 255 * 100
            except Exception:
                pass
            quality = min(100, (brightness * 0.3 + sharpness * 0.4 + min(width, height) / 40 * 0.3))
            return {
                "width": width,
                "height": height,
                "brightness": round(brightness, 1),
                "sharpness": round(sharpness, 1),
                "quality_score": round(quality, 1),
                "mode": img.mode,
                "format": img.format,
            }
    except Exception as e:
        return {"error": str(e)}


def inspect_audio(filepath: str) -> dict:
    try:
        from mutagen import File as MutagenFile
        audio = MutagenFile(filepath)
        info = {"valid": True}
        if audio:
            info["duration"] = getattr(audio.info, "length", 0)
            info["bitrate"] = getattr(audio.info, "bitrate", 0) // 1000
            info["channels"] = getattr(audio.info, "channels", 0)
            info["sample_rate"] = getattr(audio.info, "sample_rate", 0)
            if audio.tags:
                info["title"] = str(audio.tags.get("title", [""])[0]) if "title" in audio.tags else ""
                info["artist"] = str(audio.tags.get("artist", [""])[0]) if "artist" in audio.tags else ""
                info["album"] = str(audio.tags.get("album", [""])[0]) if "album" in audio.tags else ""
        return info
    except Exception as e:
        return {"valid": False, "error": str(e)}


def inspect_document(filepath: str) -> dict:
    ext = Path(filepath).suffix.lower()
    result = {"valid": True}
    try:
        if ext == ".pdf":
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(filepath)
                result["pages"] = len(reader.pages)
                if reader.metadata:
                    result["title"] = reader.metadata.get("/Title", "")
                    result["author"] = reader.metadata.get("/Author", "")
            except Exception:
                pass
        elif ext in (".docx",):
            try:
                from docx import Document
                doc = Document(filepath)
                result["paragraphs"] = len(doc.paragraphs)
            except Exception:
                pass
        elif ext in (".xlsx",):
            try:
                from openpyxl import load_workbook
                wb = load_workbook(filepath, read_only=True)
                result["sheets"] = len(wb.sheetnames)
                wb.close()
            except Exception:
                pass
    except Exception as e:
        result["error"] = str(e)
    return result


@click.command()
@click.option("--material-type", "-t", type=click.Choice(["image", "audio", "video", "document", "all"]),
              default="all", help="素材类型过滤")
@click.option("--min-quality", type=int, default=0, help="最低质量分数(0-100)")
@click.option("--show-details/--no-show-details", default=False, help="显示详细信息")
@click.option("--dry-run", is_flag=True, help="试运行")
@click.option("--fix-missing/--no-fix-missing", default=True, help="补充缺失的元数据")
@click.pass_context
def inspect_cmd(ctx, material_type, min_quality, show_details, dry_run, fix_missing):
    """检查素材格式、清晰度和元数据"""
    config: AppConfig = ctx.obj["config"]
    state: AppState = ctx.obj["state"]
    config.dry_run = dry_run or config.dry_run

    session = state.get_session()
    try:
        query = session.query(Material).filter(Material.status != ProcessStatus.FAILED)
        if material_type != "all":
            query = query.filter(Material.material_type == material_type)
        materials = query.all()
    finally:
        session.close()

    if not materials:
        console.print("[yellow]没有可检查的素材[/yellow]")
        return

    console.print(f"[green]待检查: {len(materials)} 个素材[/green]")

    task_id = state.create_batch_task(
        "inspect", f"检查 {len(materials)} 个素材",
        params={"material_type": material_type, "min_quality": min_quality}
    )

    stats = {"images": 0, "audios": 0, "videos": 0, "documents": 0, "low_quality": 0, "errors": 0}

    def process_material(material, tid):
        filepath = Path(material.current_path)
        if not filepath.exists():
            state.update_material_status(material.id, ProcessStatus.FAILED, "文件不存在")
            stats["errors"] += 1
            return False

        try:
            if material.material_type == MaterialType.IMAGE:
                info = inspect_image(str(filepath))
                if "error" not in info:
                    stats["images"] += 1
                    if not config.dry_run and fix_missing:
                        material.width = info.get("width")
                        material.height = info.get("height")
                        material.quality_score = info.get("quality_score")
                        if info.get("quality_score", 100) < min_quality:
                            stats["low_quality"] += 1
                        material.status = ProcessStatus.INSPECTED
                        from ..state import save_material
                        save_material(state, material)
                    if show_details:
                        console.print(f"  🖼️ {material.file_name}: {info.get('width')}x{info.get('height')}, "
                                      f"质量={info.get('quality_score')}%")
                else:
                    stats["errors"] += 1
                    if not config.dry_run:
                        state.update_material_status(material.id, ProcessStatus.FAILED, info.get("error"))

            elif material.material_type == MaterialType.AUDIO:
                info = inspect_audio(str(filepath))
                if info.get("valid"):
                    stats["audios"] += 1
                    if not config.dry_run and fix_missing:
                        material.duration = info.get("duration")
                        material.bitrate = info.get("bitrate")
                        if info.get("title"):
                            material.title = info.get("title")
                        if info.get("artist"):
                            material.author = info.get("artist")
                        material.status = ProcessStatus.INSPECTED
                        from ..state import save_material
                        save_material(state, material)
                    if show_details:
                        console.print(f"  🎵 {material.file_name}: "
                                      f"{human_readable_duration(info.get('duration', 0))}, "
                                      f"{info.get('bitrate', 0)}kbps")
                else:
                    stats["errors"] += 1
                    if not config.dry_run:
                        state.update_material_status(material.id, ProcessStatus.FAILED, info.get("error"))

            elif material.material_type == MaterialType.DOCUMENT:
                info = inspect_document(str(filepath))
                stats["documents"] += 1
                if not config.dry_run and fix_missing:
                    if info.get("title"):
                        material.title = info["title"]
                    if info.get("author"):
                        material.author = info["author"]
                    material.status = ProcessStatus.INSPECTED
                    from ..state import save_material
                    save_material(state, material)
                if show_details:
                    details = []
                    if "pages" in info:
                        details.append(f"{info['pages']}页")
                    if "sheets" in info:
                        details.append(f"{info['sheets']}个工作表")
                    console.print(f"  📄 {material.file_name}: {', '.join(details) if details else '已检查'}")

            elif material.material_type == MaterialType.VIDEO:
                stats["videos"] += 1
                if show_details:
                    console.print(f"  🎬 {material.file_name}: 视频(详细信息需ffmpeg支持)")
                if not config.dry_run and fix_missing:
                    material.status = ProcessStatus.INSPECTED
                    from ..state import save_material
                    save_material(state, material)

            state.log_process(material_id=material.id, task_id=tid, action="inspect", success=True)
            return True
        except Exception as e:
            stats["errors"] += 1
            console.print(f"[red]检查失败 {material.file_name}: {e}[/red]")
            return False

    result = run_batch(state, task_id, materials, process_material, description="检查素材")

    table = Table(title="检查统计")
    table.add_column("类型", style="cyan")
    table.add_column("数量", justify="right")
    table.add_row("图片", str(stats["images"]))
    table.add_row("音频", str(stats["audios"]))
    table.add_row("视频", str(stats["videos"]))
    table.add_row("文档", str(stats["documents"]))
    table.add_row("低质量", str(stats["low_quality"]))
    table.add_row("错误", str(stats["errors"]))
    console.print(table)
